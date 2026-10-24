"""Read-only anchors in actual file bytes. Location never grants claim support."""
import io
import zipfile
import xml.etree.ElementTree as ET
from workflow_io import *

NS={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
def w(name):return '{'+NS['w']+'}'+name

def checked_zip(raw):
    z=zipfile.ZipFile(io.BytesIO(raw))
    if len(z.infolist())>10000 or sum(i.file_size for i in z.infolist())>100_000_000:
        z.close();raise ValueError('OOXML archive exceeds bounded parser size')
    return z

def docx_anchor(raw,locator):
    with checked_zip(raw) as z:
        xml=z.read('word/document.xml')
    if b'<!DOCTYPE' in xml or b'<!ENTITY' in xml:raise ValueError('DTD/entities unsupported')
    root=ET.fromstring(xml);body=root.find('w:body',NS)
    if body is None:raise ValueError('DOCX body unavailable')
    headings=[];paragraphs=[];tables=[]
    def text(p):
        if p.find('.//w:txbxContent',NS) is not None or p.find('.//w:p',NS) is not None:
            raise ValueError('Nested/textbox paragraph is protected and unsupported by this locator')
        return ''.join(n.text or '' for n in p.iter(w('t')))
    for child in body:
        if child.tag==w('p'):
            content=text(child)
            style=child.find('w:pPr/w:pStyle',NS)
            st=style.get(w('val'),'') if style is not None else ''
            import re
            m=re.fullmatch(r'Heading([1-9])',st)
            if m:
                level=int(m[1]);headings=headings[:level-1]+[content]
            paragraphs.append((child,content,list(headings)))
        elif child.tag==w('tbl'):tables.append((child,list(headings)))
    kind=locator['kind']
    if kind=='paragraph':
        index=locator['paragraph']
        if type(index)!=int or not 1<=index<=len(paragraphs):raise ValueError('Paragraph does not exist')
        p,content,section=paragraphs[index-1]
        anchor={'part':'word/document.xml','body_paragraph':index,'section_path':section,
                'paragraph_text_sha256':value_sha(content),'xml_sha256':hashlib.sha256(ET.tostring(p)).hexdigest()}
    elif kind=='table_cell':
        ti,ri,ci=(locator[x] for x in ('table','row','cell'))
        if any(type(x)!=int or x<1 for x in (ti,ri,ci)) or ti>len(tables):raise ValueError('Invalid table path')
        table,section=tables[ti-1]
        if table.find('.//w:tbl',NS) is not None or table.find('.//w:vMerge',NS) is not None or table.find('.//w:gridSpan',NS) is not None:
            raise ValueError('Nested/merged table requires manual location review')
        rows=table.findall('w:tr',NS)
        if ri>len(rows):raise ValueError('Row does not exist')
        cells=rows[ri-1].findall('w:tc',NS)
        if ci>len(cells):raise ValueError('Cell does not exist')
        cell=cells[ci-1];content='\n'.join(text(p) for p in cell.findall('w:p',NS))
        anchor={'part':'word/document.xml','table':ti,'row':ri,'cell':ci,'section_path':section,'cell_text_sha256':value_sha(content)}
    else:raise ValueError('DOCX object unsupported; no invented footnote/textbox position')
    if 'section_path' in locator and locator['section_path']!=section:raise ValueError('Section path conflicts with actual document')
    return content,anchor

def locate(record,base,reader):
    if not record.get('EVID-ID'):raise ValueError('EVID-ID required')
    raw,binding=source_binding(record['source'],base,reader)
    loc=record['locator'];fmt=record['format'];quote=record.get('quote');quality='ANCHOR_CHECKED';status='Passed';text='';anchor={}
    render_sha=None;limits=[]
    if fmt in {'markdown','txt'}:
        lines=raw.decode(record.get('encoding','utf-8')).splitlines();start=loc['line_start'];end=loc.get('line_end',start)
        if type(start)!=int or type(end)!=int or not 1<=start<=end<=len(lines):raise ValueError('Invalid actual physical line range')
        text='\n'.join(lines[start-1:end]);anchor={'physical_line_start':start,'physical_line_end':end,'text_sha256':value_sha(text)}
    elif fmt=='pdf':
        if any(k in loc for k in ('line_start','line_end','physical_line','text_line')):
            raise ValueError('Extraction line numbers are not original PDF physical lines')
        import fitz
        with fitz.open(stream=raw,filetype='pdf') as doc:
            number=loc['physical_page']
            if type(number)!=int or not 1<=number<=len(doc):raise ValueError('PDF physical page does not exist')
            page=doc[number-1];label=page.get_label();rect=page.rect
            if loc.get('page_label') is not None and loc['page_label']!=label:raise ValueError('Page label conflicts with actual PDF label; printed numbers are separate')
            if 'region' in loc:
                if len(loc['region'])!=4:raise ValueError('Region needs x0,y0,x1,y1 in PDF points')
                clip=fitz.Rect(loc['region'])
                if clip.is_empty or not rect.contains(clip):raise ValueError('Region outside actual page')
            else:clip=rect
            text=page.get_text('text',clip=clip,sort=True)
            pix=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),clip=clip,alpha=False)
            render_sha=hashlib.sha256(pix.tobytes('png')).hexdigest()
            anchor={'physical_page':number,'page_label':label or None,'printed_page':'NOT_VERIFIED',
                    'page_count':len(doc),'region_points':list(clip),'coordinate_system':'PyMuPDF unrotated page points; top-left origin',
                    'page_rotation':page.rotation,'extraction_tool':'PyMuPDF '+fitz.VersionBind,'text_sha256':value_sha(text)}
            if loc.get('section') and norm(loc['section']) not in norm(page.get_text()):raise ValueError('Requested section not found on stated physical page')
            quality='EXTRACTED_NOT_VISUALLY_CHECKED' if text.strip() else 'REGION_ONLY_OCR_NOT_RUN'
            status='Partial';limits=['Text layer does not prove visible page content. No automatic OCR executed.']
    elif fmt=='image':
        from PIL import Image
        with Image.open(io.BytesIO(raw)) as im:
            region=loc['region_pixels']
            if len(region)!=4 or any(type(x)!=int for x in region) or not 0<=region[0]<region[2]<=im.width or not 0<=region[1]<region[3]<=im.height:raise ValueError('Invalid pixel region')
            crop=im.crop(region).convert('RGB');render_sha=hashlib.sha256(crop.tobytes()).hexdigest()
            anchor={'region_pixels':region,'image_size':[im.width,im.height],'coordinates':'top-left pixel origin','render_hash_kind':'RGB pixels'}
        quality='REGION_ONLY_OCR_NOT_RUN';status='Partial';limits=['No automatic OCR; any supplied transcription needs visual confirmation.']
    elif fmt=='docx':
        if any(k in loc for k in ('physical_page','line_start','line_end')):raise ValueError('DOCX has no fixed physical lines/pages without verified rendering')
        text,anchor=docx_anchor(raw,loc);limits=['Main-body paragraphs/simple tables only. No page layout/footnote/textbox/field evaluation claim.']
    elif fmt=='xlsx':
        import openpyxl
        from openpyxl.utils.cell import range_boundaries
        with checked_zip(raw):pass
        formulas=openpyxl.load_workbook(io.BytesIO(raw),data_only=False,read_only=True,keep_links=False)
        cached=openpyxl.load_workbook(io.BytesIO(raw),data_only=True,read_only=True,keep_links=False)
        try:
            sheet=loc['sheet'];cellrange=loc['range']
            if sheet not in formulas.sheetnames:raise ValueError('Worksheet does not exist')
            minc,minr,maxc,maxr=range_boundaries(cellrange)
            if None in (minc,minr,maxc,maxr) or min(minc,minr)<1 or maxc<minc or maxr<minr or (maxc-minc+1)*(maxr-minr+1)>100:raise ValueError('Use a bounded <=100-cell range')
            cells={}
            for row in formulas[sheet].iter_rows(min_row=minr,max_row=maxr,min_col=minc,max_col=maxc):
                for cell in row:
                    cells[cell.coordinate]={'formula':cell.value if cell.data_type=='f' else None,'value':None if cell.data_type=='f' else cell.value,
                                            'cached_value':cached[sheet][cell.coordinate].value if cell.data_type=='f' else None,'number_format':cell.number_format}
            headers={k:formulas[sheet][address].value for k,address in loc.get('headers',{}).items()}
            for address,expect in record.get('expected_cells',{}).items():
                if address not in cells or any(cells[address].get(k)!=v for k,v in expect.items()):raise ValueError('Cell content/formula differs from requested anchor')
            anchor={'sheet':sheet,'range':cellrange,'cells':cells,'headers':headers,'recalculated':False,'reader':'openpyxl '+openpyxl.__version__}
            text=json.dumps(cells,ensure_ascii=False,default=str)
            limits=['Formula and cached value are distinct. Cached values may be absent/stale; no recalculation.']
        finally:formulas.close();cached.close()
    else:raise ValueError('Unsupported file format; manual location required')
    if quote:
        if text and norm(quote) not in norm(text):raise ValueError('Quote absent at actual location')
        if not text and fmt not in {'pdf','image'}:raise ValueError('No text at anchor')
    elif fmt not in {'pdf','image','xlsx'}:raise ValueError('Textual anchor needs a nonempty excerpt')
    binding_context={'EVID-ID':record['EVID-ID'],'source':binding,'locator':loc,'actual_anchor':anchor,'quote':quote,'render_sha256':render_sha}
    fingerprint=value_sha(stable(binding_context));visual='NOT_SUPPLIED'
    if record.get('visual_review_ref'):
        review=reader.evidence(record['visual_review_ref'],base)
        if not render_sha or review.get('context_sha256')!=fingerprint or review.get('render_sha256')!=render_sha or not all(review.get(k) for k in ('reviewer','reviewed_at','scope','reason')) or type(review.get('matches'))!=bool:
            raise ValueError('Visual review stale or not bound to actual page rendering')
        timestamp(review['reviewed_at']);visual='CURRENT_SUPPLIED_REVIEW'
        if review['matches']:
            status='Passed';quality='VISUALLY_CONFIRMED_BY_SUPPLIED_REVIEW'
        else:status='Failed';quality='VISUAL_CONFLICT';limits.append('Supplied actual visual review rejects the proposed text/anchor')
    return {'EVID-ID':record['EVID-ID'],'SRC-ID':record['source']['SRC-ID'],'status':status,'location_quality':quality,
            'anchor':anchor,'excerpt':text,'requested_quote':quote,'source_binding':binding,'context_sha256':fingerprint,
            'render_sha256':render_sha,'visual_review':visual,'claim_support':'Unknown','paper_gate':'NOT RUN','errors':[],
            'limits':limits+['Source identity and semantic support remain independent. Saved reviewer identity/authenticity not certified.']}

if __name__=='__main__':entry(locate,__doc__)
